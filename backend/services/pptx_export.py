"""
PPTX Export Servisi — v2 (Profesyonel Tasarım)
python-pptx ile Pitch ve Yatırımcı raporları üretir.

Düzeltmeler (v2):
  - Kapak başlığı: ham fikir cümlesi yerine profesyonel bir başlık
  - İçerik: çift bullet sorunları ve ham markdown (**, ---, •) temizlendi
  - Font boyutları artırıldı
  - Yatırımcı raporu: bölüm başlıkları ayrı stilde gösteriliyor
"""

from __future__ import annotations

import io
import re
import textwrap
from typing import List, Tuple

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# ─── Renk paleti ─────────────────────────────────────────────────────────────
COLOR_TEAL       = RGBColor(0x1A, 0x7A, 0x6E)
COLOR_TEAL_DARK  = RGBColor(0x10, 0x5C, 0x57)
COLOR_TEAL_LIGHT = RGBColor(0xB2, 0xD8, 0xD4)
COLOR_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_DARK       = RGBColor(0x1C, 0x27, 0x23)
COLOR_LIGHT_BG   = RGBColor(0xF4, 0xF7, 0xF6)
COLOR_ACCENT     = RGBColor(0xE8, 0x6C, 0x4B)   # başlık vurgu çizgisi
COLOR_MUTED      = RGBColor(0x5A, 0x6B, 0x66)

# ─── Sayfa boyutları (16:9 geniş) ─────────────────────────────────────────────
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# ─── Font boyutları ───────────────────────────────────────────────────────────
COVER_TITLE_FONT   = Pt(36)
COVER_SUBTITLE_FONT = Pt(18)
SLIDE_TITLE_FONT   = Pt(30)
BODY_FONT          = Pt(17)
SECTION_FONT       = Pt(19)
SMALL_FONT         = Pt(13)


# ─── Yardımcı: Profesyonel başlık üretici ────────────────────────────────────

def _generate_cover_title(idea: str) -> str:
    """
    Ham fikir cümlesinden profesyonel, kısa bir başlık üretir.
    'geliştirmek istiyorum', 'yapmak istiyorum' gibi sonekleri temizler.
    """
    title = idea.strip().rstrip('.')

    # Yaygın Türkçe fiil + niyet yapılarını sil
    patterns = [
        r'\s+(geliştirmek|yapmak|kurmak|oluşturmak|hazırlamak|üretmek|'
        r'tasarlamak|inşa etmek|başlatmak)\s+istiyorum\.?',
        r'\s+(geliştirmeyi|yapmayı|kurmayı|oluşturmayı|hazırlamayı)\s+'
        r'(planlıyorum|düşünüyorum|hedefliyorum)\.?',
        r'\s+istiyorum\.?',
        r'\s+planlıyorum\.?',
        r'\s+hedefliyorum\.?',
    ]
    for pattern in patterns:
        title = re.sub(pattern, '', title, flags=re.IGNORECASE).strip()

    title = title.rstrip('.,;:').strip()

    # İlk harf büyük
    if title:
        title = title[0].upper() + title[1:]

    # 65 karakterden uzunsa kısalt
    if len(title) > 65:
        words = title.split()
        short = ""
        for w in words:
            if len(short) + len(w) + 1 <= 65:
                short = (short + " " + w).strip()
            else:
                break
        title = short.rstrip('.,;:') + "…"

    return title


# ─── Yardımcı: Inline Markdown temizleyici ────────────────────────────────────

def _strip_inline_md(text: str) -> str:
    """**bold**, *italic*, `code`, [link](url) gibi inline markdown'ı temizler."""
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)   # ***bold italic***
    text = re.sub(r'\*\*(.+?)\*\*',     r'\1', text)   # **bold**
    text = re.sub(r'\*(.+?)\*',         r'\1', text)   # *italic*
    text = re.sub(r'__(.+?)__',         r'\1', text)   # __bold__
    text = re.sub(r'_(.+?)_',           r'\1', text)   # _italic_
    text = re.sub(r'`(.+?)`',           r'\1', text)   # `code`
    text = re.sub(r'\[(.+?)\]\(.+?\)',  r'\1', text)   # [text](url)
    return text.strip()


# ─── Yardımcı: Markdown satır ayrıştırıcı ────────────────────────────────────

def _is_section_header(line: str) -> bool:
    """Satırın bölüm başlığı olup olmadığını kontrol eder."""
    # Markdown başlık (#, ##, ###)
    if re.match(r'^#{1,3}\s+', line):
        return True
    # Numaralı büyük başlık (1. HANGI TÜR...)
    if re.match(r'^\d+[\.\)]\s+[A-ZÇĞİÖŞÜ]', line):
        return True
    # ▸ ile başlayan büyük harf başlık
    if re.match(r'^▸\s+[A-ZÇĞİÖŞÜ]', line):
        return True
    return False


def _parse_structured_lines(text: str) -> List[Tuple[str, str]]:
    """
    Markdown metnini (tip, içerik) tuple listesi olarak döndürür.
    tip: 'header' | 'bullet' | 'text' | 'skip'
    """
    result = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        # Yatay çizgi → atla
        if re.match(r'^[-_*]{3,}$', line):
            result.append(('skip', ''))
            continue
        # Markdown başlık
        if re.match(r'^#{1,3}\s+', line):
            clean = re.sub(r'^#{1,3}\s+', '', line).strip()
            result.append(('header', _strip_inline_md(clean)))
            continue
        # Numaralı başlık (1. BAŞLIK, 2. BAŞLIK)
        m = re.match(r'^(\d+[\.\)])\s+(.+)$', line)
        if m and m.group(2) == m.group(2).upper() and len(m.group(2)) > 3:
            result.append(('header', _strip_inline_md(f"{m.group(1)} {m.group(2)}")))
            continue
        # Bullet: -, *, +, •, ▸, →
        if re.match(r'^[-*+•▸→]\s+', line):
            content = re.sub(r'^[-*+•▸→]\s+', '', line).strip()
            result.append(('bullet', _strip_inline_md(content)))
            continue
        # Numaralı madde (1. küçük metin)
        m2 = re.match(r'^\d+[\.\)]\s+(.+)$', line)
        if m2:
            result.append(('bullet', _strip_inline_md(m2.group(1))))
            continue
        # Düz metin
        result.append(('text', _strip_inline_md(line)))
    return [(t, c) for t, c in result if t != 'skip' and c]


# ─── Yardımcı: Taşma hesaplama ───────────────────────────────────────────────

def _line_count(text: str, chars_per_line: int = 80) -> float:
    wrapped = textwrap.wrap(text, width=chars_per_line) or [""]
    return float(len(wrapped))


def _chunk_structured(
    lines: List[Tuple[str, str]],
    max_lines: float = 14.0,
    chars_per_line: int = 80,
) -> List[List[Tuple[str, str]]]:
    """Yapılandırılmış satır listesini slayt chunk'larına böler."""
    chunks: List[List[Tuple[str, str]]] = []
    current: List[Tuple[str, str]] = []
    current_h = 0.0

    for typ, content in lines:
        h = _line_count(content, chars_per_line) + (0.5 if typ == 'header' else 0)
        if current and current_h + h > max_lines:
            chunks.append(current)
            current = []
            current_h = 0.0
        current.append((typ, content))
        current_h += h

    if current:
        chunks.append(current)
    return chunks if chunks else [[]]


# ─── Slayt oluşturma yardımcıları ────────────────────────────────────────────

def _apply_bg(slide, color: RGBColor) -> None:
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_title_bar(slide, title_text: str) -> None:
    """Slayt üstüne teal başlık şeridi ve beyaz başlık metni ekler."""
    # Arka plan şerit
    bar = slide.shapes.add_shape(
        1,  # RECTANGLE
        left=0, top=0,
        width=SLIDE_W, height=Inches(1.4),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_TEAL_DARK
    bar.line.fill.background()

    # Alt vurgu çizgisi (turuncu)
    accent = slide.shapes.add_shape(
        1,
        left=0, top=Inches(1.4),
        width=SLIDE_W, height=Inches(0.06),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ACCENT
    accent.line.fill.background()

    # Başlık metni
    txBox = slide.shapes.add_textbox(
        left=Inches(0.7), top=Inches(0.1),
        width=Inches(12.0), height=Inches(1.2),
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title_text
    run.font.bold = True
    run.font.size = SLIDE_TITLE_FONT
    run.font.color.rgb = COLOR_WHITE


def _add_structured_content(slide, lines: List[Tuple[str, str]]) -> None:
    """Yapılandırılmış içeriği içerik alanına ekler (başlık/bullet ayrı stilde)."""
    txBox = slide.shapes.add_textbox(
        left=Inches(0.7), top=Inches(1.65),
        width=Inches(11.9), height=Inches(5.6),
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None

    first = True
    for typ, content in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()

        if typ == 'header':
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(10)
            run = p.add_run()
            run.text = content
            run.font.bold = True
            run.font.size = SECTION_FONT
            run.font.color.rgb = COLOR_TEAL
        elif typ == 'bullet':
            p.alignment = PP_ALIGN.LEFT
            p.level = 1
            run = p.add_run()
            run.text = f"  ▸  {content}"
            run.font.size = BODY_FONT
            run.font.color.rgb = COLOR_DARK
        else:  # 'text'
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = content
            run.font.size = BODY_FONT
            run.font.color.rgb = COLOR_MUTED


def _add_slide_number(slide, number: int, total: int) -> None:
    """Sağ alt köşeye slayt numarası ekler."""
    txBox = slide.shapes.add_textbox(
        left=Inches(11.8), top=Inches(7.05),
        width=Inches(1.3), height=Inches(0.35),
    )
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{number} / {total}"
    run.font.size = SMALL_FONT
    run.font.color.rgb = COLOR_MUTED


# ─── Sunum Oluşturucular ─────────────────────────────────────────────────────

def build_pitch_pptx(idea: str, slides: list) -> bytes:
    """Pitch deck için profesyonel PPTX üretir."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    professional_title = _generate_cover_title(idea)
    total_slides = 1 + sum(
        len(_chunk_structured(_parse_structured_lines(s.get("content", ""))))
        for s in slides
    ) + 1

    # ── Kapak Slaytı ──────────────────────────────────────────────────────────
    cover = prs.slides.add_slide(blank)
    _apply_bg(cover, COLOR_TEAL_DARK)

    # Dekoratif üst şerit
    top_bar = cover.shapes.add_shape(
        1, left=0, top=0, width=SLIDE_W, height=Inches(0.18)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT
    top_bar.line.fill.background()

    # Ana başlık
    txMain = cover.shapes.add_textbox(
        left=Inches(1.2), top=Inches(1.8),
        width=Inches(10.9), height=Inches(2.5),
    )
    tf = txMain.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = professional_title
    run.font.bold = True
    run.font.size = COVER_TITLE_FONT
    run.font.color.rgb = COLOR_WHITE

    # Alt çizgi / ayraç
    sep = cover.shapes.add_shape(
        1, left=Inches(5.2), top=Inches(4.5),
        width=Inches(2.9), height=Inches(0.06)
    )
    sep.fill.solid()
    sep.fill.fore_color.rgb = COLOR_TEAL_LIGHT
    sep.line.fill.background()

    # Etiket (Pitch Deck)
    txSub = cover.shapes.add_textbox(
        left=Inches(1.2), top=Inches(4.7),
        width=Inches(10.9), height=Inches(0.8),
    )
    tf2 = txSub.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = "PITCH DECK"
    run2.font.size = COVER_SUBTITLE_FONT
    run2.font.bold = True
    run2.font.color.rgb = COLOR_TEAL_LIGHT

    # ── İçerik Slaytları ──────────────────────────────────────────────────────
    slide_num = 2
    for slide_data in slides:
        title = slide_data.get("title", "")
        content_text = slide_data.get("content", "")

        structured = _parse_structured_lines(content_text)
        # Eğer hiç başlık/bullet yoksa düz metin olarak işle
        if not structured:
            structured = [('text', content_text.strip())]

        chunks = _chunk_structured(structured)
        for chunk_idx, chunk in enumerate(chunks):
            sl = prs.slides.add_slide(blank)
            _apply_bg(sl, COLOR_LIGHT_BG)
            slide_title = title if chunk_idx == 0 else f"{title} (devam)"
            _add_title_bar(sl, slide_title)
            _add_structured_content(sl, chunk)
            _add_slide_number(sl, slide_num, total_slides)
            slide_num += 1

    # ── Kapanış Slaytı ────────────────────────────────────────────────────────
    closing = prs.slides.add_slide(blank)
    _apply_bg(closing, COLOR_TEAL_DARK)

    top_bar2 = closing.shapes.add_shape(
        1, left=0, top=0, width=SLIDE_W, height=Inches(0.18)
    )
    top_bar2.fill.solid()
    top_bar2.fill.fore_color.rgb = COLOR_ACCENT
    top_bar2.line.fill.background()

    txClose = closing.shapes.add_textbox(
        left=Inches(1.5), top=Inches(2.4),
        width=Inches(10.3), height=Inches(1.8),
    )
    tf3 = txClose.text_frame
    p3 = tf3.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    run3 = p3.add_run()
    run3.text = "Teşekkürler"
    run3.font.bold = True
    run3.font.size = Pt(44)
    run3.font.color.rgb = COLOR_WHITE

    p4 = tf3.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    run4 = p4.add_run()
    run4.text = professional_title
    run4.font.size = Pt(18)
    run4.font.color.rgb = COLOR_TEAL_LIGHT

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()


def build_investors_pptx(idea: str, advice: str) -> bytes:
    """Yatırımcı analiz raporu için profesyonel PPTX üretir."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    professional_title = _generate_cover_title(idea)
    structured = _parse_structured_lines(advice)
    if not structured:
        structured = [('text', advice.strip())]

    chunks = _chunk_structured(structured)
    total_slides = 1 + len(chunks)

    # ── Kapak Slaytı ──────────────────────────────────────────────────────────
    cover = prs.slides.add_slide(blank)
    _apply_bg(cover, COLOR_TEAL_DARK)

    top_bar = cover.shapes.add_shape(
        1, left=0, top=0, width=SLIDE_W, height=Inches(0.18)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT
    top_bar.line.fill.background()

    txMain = cover.shapes.add_textbox(
        left=Inches(1.2), top=Inches(1.6),
        width=Inches(10.9), height=Inches(2.8),
    )
    tf = txMain.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = professional_title
    run.font.bold = True
    run.font.size = COVER_TITLE_FONT
    run.font.color.rgb = COLOR_WHITE

    sep = cover.shapes.add_shape(
        1, left=Inches(5.2), top=Inches(4.5),
        width=Inches(2.9), height=Inches(0.06)
    )
    sep.fill.solid()
    sep.fill.fore_color.rgb = COLOR_TEAL_LIGHT
    sep.line.fill.background()

    txSub = cover.shapes.add_textbox(
        left=Inches(1.2), top=Inches(4.7),
        width=Inches(10.9), height=Inches(0.8),
    )
    tf2 = txSub.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = "YATIRIMCI BULMA RAPORU"
    run2.font.size = COVER_SUBTITLE_FONT
    run2.font.bold = True
    run2.font.color.rgb = COLOR_TEAL_LIGHT

    # ── İçerik Slaytları ──────────────────────────────────────────────────────
    for chunk_idx, chunk in enumerate(chunks):
        sl = prs.slides.add_slide(blank)
        _apply_bg(sl, COLOR_LIGHT_BG)
        title = "Yatırımcı Tavsiyesi" if chunk_idx == 0 else f"Yatırımcı Tavsiyesi — Devam {chunk_idx + 1}"
        _add_title_bar(sl, title)
        _add_structured_content(sl, chunk)
        _add_slide_number(sl, chunk_idx + 2, total_slides)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
